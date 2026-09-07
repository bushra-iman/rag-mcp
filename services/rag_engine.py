import os
import re
from config import Config
from langchain_openai import (
    ChatOpenAI,
    OpenAIEmbeddings
)
from langchain_core.documents import (
    Document
)
from langchain_community.vectorstores import (
    FAISS
)
from langchain_text_splitters import (
    RecursiveCharacterTextSplitter
)
from langchain.agents import (
    create_agent
)
from pypdf import PdfReader
from docx import Document as DocxDocument
from services.web_search import (
    search_web
)
from services.agent_memory import (
    memory,
    list_threads,
    get_top_conversations,
    build_thread_id
)
from mcp_integration.client_manager import (
    mcp_client_manager
)
from mcp_integration.tool_registry import (
    mcp_tool_registry
)
from mcp_integration.tool_adapter import (
    create_mcp_langchain_tool
)
# =========================================================
# CONFIGURATION
# =========================================================
llm = ChatOpenAI(
    api_key=Config.OPENAI_API_KEY,
    model="gpt-4.1-mini",
    temperature=0
)
embeddings = OpenAIEmbeddings(
    api_key=Config.OPENAI_API_KEY
)
VECTOR_DB = (
    Config.VECTOR_DB_PATH
    or "vector_store"
)
TEXT_INDEX = (
    "document_index.txt"
)
MAX_FILES = 10
MAX_FILE_SIZE = (
    10 * 1024 * 1024
)
ALLOWED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt"
}
# =========================================================
# FILE VALIDATION
# =========================================================
def validate_file(file):
    if not file or not file.filename:
        raise ValueError(
            "Invalid file."
        )
    extension = os.path.splitext(
        file.filename
    )[1].lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError(
            "Unsupported file type. "
            "Only PDF, DOCX and TXT files are allowed."
        )
    file.seek(
        0,
        os.SEEK_END
    )
    file_size = file.tell()
    file.seek(0)
    if file_size > MAX_FILE_SIZE:
        raise ValueError(
            "File is too large. "
            "Maximum allowed size is 10 MB per file."
        )
    return extension
# =========================================================
# PDF TEXT EXTRACTION
# =========================================================
def extract_pdf_text(file):
    reader = PdfReader(
        file
    )
    text = ""
    for page in reader.pages:
        page_text = (
            page.extract_text()
        )
        if page_text:
            text += (
                page_text
                + "\n"
            )
    return text
# =========================================================
# DOCX TEXT EXTRACTION
# =========================================================
def extract_docx_text(file):
    document = DocxDocument(
        file
    )
    text = ""
    for paragraph in (
        document.paragraphs
    ):
        if paragraph.text.strip():
            text += (
                paragraph.text
                + "\n"
            )
    return text
# =========================================================
# TXT TEXT EXTRACTION
# =========================================================
def extract_txt_text(file):
    content = file.read()
    return content.decode(
        "utf-8",
        errors="ignore"
    )
# =========================================================
# EXTRACT TEXT
# =========================================================
def extract_text_from_file(
    file
):
    extension = validate_file(
        file
    )
    if extension == ".pdf":
        return extract_pdf_text(
            file
        )
    if extension == ".docx":
        return extract_docx_text(
            file
        )
    if extension == ".txt":
        return extract_txt_text(
            file
        )
    raise ValueError(
        "Unsupported file type."
    )
# =========================================================
# CREATE / UPDATE VECTOR STORE
# =========================================================
def create_vector_store_from_files(
    files
):
    if not files:
        raise ValueError(
            "No files were uploaded."
        )
    if len(files) > MAX_FILES:
        raise ValueError(
            f"Maximum {MAX_FILES} files "
            "can be uploaded in one request."
        )
    documents = []
    uploaded_files = []
    for file in files:
        text = extract_text_from_file(
            file
        )
        if not text.strip():
            continue
        documents.append(
            Document(
                page_content=text,
                metadata={
                    "source":
                        file.filename
                }
            )
        )
        uploaded_files.append(
            file.filename
        )
    if not documents:
        raise ValueError(
            "No readable text was found "
            "in the uploaded files."
        )
    # -----------------------------------------------------
    # CHUNKING
    # -----------------------------------------------------
    splitter = (
        RecursiveCharacterTextSplitter(
            chunk_size=500,
            chunk_overlap=50
        )
    )
    chunks = (
        splitter.split_documents(
            documents
        )
    )
    # -----------------------------------------------------
    # FAISS
    # -----------------------------------------------------
    if os.path.exists(
        VECTOR_DB
    ):
        db = FAISS.load_local(
            VECTOR_DB,
            embeddings,
            allow_dangerous_deserialization=True
        )
        db.add_documents(
            chunks
        )
    else:
        db = FAISS.from_documents(
            chunks,
            embeddings
        )
    db.save_local(
        VECTOR_DB
    )
    # -----------------------------------------------------
    # EXACT TEXT INDEX
    # -----------------------------------------------------
    with open(
        TEXT_INDEX,
        "a",
        encoding="utf-8"
    ) as index_file:
        for document in documents:
            index_file.write(
                "\n"
            )
            index_file.write(
                "=" * 80
            )
            index_file.write(
                "\n"
            )
            index_file.write(
                "SOURCE: "
                f"{document.metadata.get('source', 'unknown')}\n"
            )
            index_file.write(
                document.page_content
            )
            index_file.write(
                "\n"
            )
    return {
        "message":
            "Files uploaded and indexed successfully.",
        "files_uploaded":
            uploaded_files,
        "total_files":
            len(uploaded_files),
        "chunks_created":
            len(chunks)
    }
# =========================================================
# OLD TEXT INGESTION
# =========================================================
def create_vector_store(
    text
):
    if not text or not text.strip():
        raise ValueError(
            "Text cannot be empty."
        )
    splitter = (
        RecursiveCharacterTextSplitter(
            chunk_size=500,
            chunk_overlap=50
        )
    )
    docs = [
        Document(
            page_content=text,
            metadata={
                "source":
                    "manual_text"
            }
        )
    ]
    chunks = (
        splitter.split_documents(
            docs
        )
    )
    if os.path.exists(
        VECTOR_DB
    ):
        db = FAISS.load_local(
            VECTOR_DB,
            embeddings,
            allow_dangerous_deserialization=True
        )
        db.add_documents(
            chunks
        )
    else:

        db = FAISS.from_documents(
            chunks,
            embeddings
        )
    db.save_local(
        VECTOR_DB
    )
    return (
        "Document added successfully."
    )
# =========================================================
# RETRIEVAL
# =========================================================
def retrieve(
    question
):
    if not os.path.exists(
        VECTOR_DB
    ):
        return []
    db = FAISS.load_local(
        VECTOR_DB,
        embeddings,
        allow_dangerous_deserialization=True
    )
    # -----------------------------------------------------
    # SEMANTIC SEARCH
    # -----------------------------------------------------
    docs = db.similarity_search(
        question,
        k=10
    )
    # -----------------------------------------------------
    # EXACT WMD ISSUE ID SEARCH
    # -----------------------------------------------------
    issue_ids = re.findall(
        r"\bwmd-\d+\b",
        question.lower()
    )
    if (
        issue_ids
        and os.path.exists(
            TEXT_INDEX
        )
    ):
        with open(
            TEXT_INDEX,
            "r",
            encoding="utf-8"
        ) as index_file:
            text = index_file.read()
        paragraphs = text.split(
            "=" * 80
        )
        exact_matches = []
        for paragraph in paragraphs:
            paragraph_lower = (
                paragraph.lower()
            )
            if any(
                issue_id in paragraph_lower
                for issue_id in issue_ids
            ):
                exact_matches.append(
                    Document(
                        page_content=paragraph,
                        metadata={
                            "source":
                                "Uploaded Document"
                        }
                    )
                )
        if exact_matches:
            return exact_matches[:10]
    return docs[:8]
# =========================================================
# LOCAL RAG TOOL
# =========================================================
def search_knowledge_base(
    question: str
) -> str:
    """
    Search uploaded documents.
    """
    docs = retrieve(
        question
    )
    if not docs:
        return (
            "NO_KNOWLEDGE_BASE_RESULT"
        )
    context = "\n\n".join(
        [
            doc.page_content
            for doc in docs
        ]
    )
    return context
# =========================================================
# LOCAL WEB SEARCH TOOL
# =========================================================
def search_web_tool(
    question: str
) -> str:
    """
    Search the web when necessary.
    """
    result = search_web(
        question
    )
    if not result:
        return (
            "NO_WEB_RESULT"
        )
    return result
# =========================================================
# BUILD ALL AGENT TOOLS
# =========================================================
def get_all_agent_tools():
    # -----------------------------------------------------
    # EXISTING LOCAL TOOLS
    # -----------------------------------------------------
    tools = [
        search_knowledge_base,
        search_web_tool
    ]
    # -----------------------------------------------------
    # DYNAMIC MCP TOOLS
    # -----------------------------------------------------
    for tool_definition in (
        mcp_tool_registry.get_all_tools()
    ):
        mcp_tool = (
            create_mcp_langchain_tool(
                mcp_client_manager,
                tool_definition
            )
        )
        tools.append(
            mcp_tool
        )
    return tools
# =========================================================
# BUILD AGENT
# =========================================================
def build_agent():
    all_tools = (
        get_all_agent_tools()
    )
    return create_agent(
        model=llm,
        tools=all_tools,
        checkpointer=memory,
        system_prompt="""
You are a Student Assistant AI.
You have access to two types of tools.
LOCAL TOOLS:
1. search_knowledge_base
   - Searches documents uploaded by the user.
2. search_web_tool
   - Searches the web.
MCP TOOLS:
- MCP tools are dynamically discovered from
  registered external MCP servers.
- Their exposed names normally start with
  "mcp_".
- Do not assume what MCP tools exist.
- Use the tool description and input schema
  to understand each MCP tool.
IMPORTANT RULES:
1. Prefer uploaded document information when
   the answer exists in the uploaded documents.
2. Use web search when uploaded documents do
   not contain enough information.
3. MCP tools should be selected automatically
   when they are appropriate for the user's request.
4. The user does not need to specify an MCP
   server manually.
5. Treat MCP tools like normal available tools.
6. Never invent information.
7. Use previous conversation messages when
   the user refers to earlier discussion.
8. Keep answers clear and concise.
"""
    )
# =========================================================
# ASK QUESTION
# =========================================================
def ask_question(
    question,
    thread_id="default",
    tenant_id="default-tenant"
):
    # -----------------------------------------------------
    # TENANT-SPECIFIC THREAD
    # -----------------------------------------------------
    unique_thread_id = (
        build_thread_id(
            tenant_id,
            thread_id
        )
    )
    # -----------------------------------------------------
    # CONFIGURATION
    # -----------------------------------------------------
    config = {
        "configurable": {
            "thread_id":
                unique_thread_id
        }
    }
    # -----------------------------------------------------
    # IMPORTANT:
    # BUILD A FRESH AGENT

    # This ensures newly registered MCP tools
    # become available immediately.
    # -----------------------------------------------------
    agent = build_agent()
    # -----------------------------------------------------
    # INVOKE AGENT
    # -----------------------------------------------------
    result = agent.invoke(
        {
            "messages": [
                {
                    "role":
                        "user",
                    "content":
                        question
                }
            ]
        },
        config
    )
    # -----------------------------------------------------
    # FINAL ANSWER
    # -----------------------------------------------------
    messages = result.get(
        "messages",
        []
    )
    if not messages:
        answer = (
            "No response was generated."
        )
    else:
        answer = (
            messages[-1].content
        )
    # -----------------------------------------------------
    # DETERMINE SOURCES
    # -----------------------------------------------------
    sources = []
    for message in messages:
        tool_name = getattr(
            message,
            "name",
            None
        )
        if not tool_name:
            continue
        if tool_name == (
            "search_knowledge_base"
        ):
            if (
                "Student Knowledge Base"
                not in sources
            ):
                sources.append(
                    "Student Knowledge Base"
                )
        elif tool_name == (
            "search_web_tool"
        ):
            if (
                "DuckDuckGo Web Search"
                not in sources
            ):
                sources.append(
                    "DuckDuckGo Web Search"
                )
        elif tool_name.startswith(
            "mcp_"
        ):
            tool_definition = (
                mcp_tool_registry
                .get_tool(
                    tool_name
                )
            )

            if tool_definition:
                original_tool_name = (
                    tool_definition.get(
                        "original_name"
                    )
                )

                if original_tool_name:
                    source = (
                        "MCP Tool: "
                        f"{original_tool_name}"
                    )
                else:
                    source = (
                        "MCP Tool: "
                        f"{tool_name}"
                    )

            else:
                source = (
                    "MCP Tool: "
                    f"{tool_name}"
                )

            if source not in sources:
                sources.append(
                    source
                )

                    # -----------------------------------------------------
    # RETURN FINAL RESPONSE
    # -----------------------------------------------------

    return {
        "answer": answer,
        "sources": sources,
        "tenant_id": tenant_id,
        "thread_id": thread_id
    }
# =========================================================
# LIST MEMORY THREADS
# =========================================================
def get_memory_threads(
    tenant_id
):
    return list_threads(
        tenant_id
    )
# =========================================================
# TOP 5 CONVERSATIONS
# =========================================================
def get_top_5_conversations(
    tenant_id
):
    return get_top_conversations(
        tenant_id,
        limit=5
    )
